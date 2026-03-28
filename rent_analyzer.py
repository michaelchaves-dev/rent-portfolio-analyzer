"""
=============================================================
  RENT PORTFOLIO ANALYZER — by Michael Chaves
  Analyzes RentManager CSV export to detect:
    1. Tenants paying below current market/lease rate
    2. Vacant units generating zero income
    3. Total monthly & annual revenue leakage
=============================================================

HOW TO EXPORT FROM RENTMANAGER:
  1. Go to Reports > Residents/Tenants > Tenant Listing
  2. Click "Export to Excel" or "Export to CSV"
  3. Save the file and note the file path
  4. Update INPUT_FILE below to match your file path

REQUIRED COLUMNS IN YOUR EXPORT (column names may vary slightly):
  - Unit (unit number)
  - Tenant Name (or First Name + Last Name)
  - Lease Start (lease start date)
  - Lease End (lease end date)
  - Market Rent (the current correct/market rate for the unit)
  - Actual Rent (what the tenant is currently paying)
  - Status (occupied / vacant)

=============================================================
"""

import pandas as pd
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter
from datetime import datetime
import os
import sys

# ─────────────────────────────────────────────
#  CONFIG — UPDATE THESE TO MATCH YOUR FILE
# ─────────────────────────────────────────────
INPUT_FILE = "tenant_export.csv"   # Your RentManager export file (CSV or XLSX)
OUTPUT_FILE = "Portfolio_Revenue_Analysis.xlsx"

# Column name mapping — update if your export uses different column names
# Left side = what this script expects | Right side = your actual column header
COLUMN_MAP = {
    "unit":         "Unit",           # Unit number/identifier
    "tenant_name":  "Tenant Name",    # Full tenant name (or adjust below for split names)
    "lease_start":  "Lease Start",    # Lease start date
    "lease_end":    "Lease End",      # Lease end date
    "market_rent":  "Market Rent",    # Current market/correct rent rate
    "actual_rent":  "Actual Rent",    # What tenant is currently paying
    "status":       "Status",         # Occupied / Vacant
    "property":     "Property",       # Property name (if multi-property)
}

# ─────────────────────────────────────────────
#  COLORS
# ─────────────────────────────────────────────
NAVY        = "1B2A4A"
ACCENT_BLUE = "2563EB"
RED         = "DC2626"
GREEN       = "16A34A"
YELLOW      = "FEF08A"
LIGHT_BLUE  = "EFF6FF"
LIGHT_RED   = "FEF2F2"
LIGHT_GREEN = "F0FDF4"
GRAY        = "F8FAFC"
WHITE       = "FFFFFF"
MID_GRAY    = "94A3B8"

def load_data(filepath):
    """Load CSV or Excel export from RentManager."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".csv":
        df = pd.read_csv(filepath, dtype=str)
    elif ext in [".xlsx", ".xls"]:
        df = pd.read_excel(filepath, dtype=str)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
    df.columns = df.columns.str.strip()
    return df

def normalize_columns(df):
    """Rename columns to standard names based on COLUMN_MAP."""
    reverse_map = {v: k for k, v in COLUMN_MAP.items()}
    df = df.rename(columns=reverse_map)
    
    # Handle split first/last name columns if no combined "Tenant Name"
    if "tenant_name" not in df.columns:
        if "First Name" in df.columns and "Last Name" in df.columns:
            df["tenant_name"] = df["First Name"].fillna("") + " " + df["Last Name"].fillna("")
            df["tenant_name"] = df["tenant_name"].str.strip()
        else:
            df["tenant_name"] = "Unknown"
    
    # Fill missing property column
    if "property" not in df.columns:
        df["property"] = "Portfolio"
    
    return df

def clean_currency(series):
    """Strip $, commas, spaces and convert to float."""
    return pd.to_numeric(
        series.astype(str)
              .str.replace(r'[$,\s]', '', regex=True)
              .str.replace(r'[()]', '', regex=True),
        errors='coerce'
    ).fillna(0.0)

def analyze(df):
    """Core analysis — flag underpaying tenants and vacant units."""
    df["market_rent_num"] = clean_currency(df.get("market_rent", pd.Series(["0"]*len(df))))
    df["actual_rent_num"] = clean_currency(df.get("actual_rent", pd.Series(["0"]*len(df))))
    
    # Determine vacancy
    if "status" in df.columns:
        df["is_vacant"] = df["status"].str.strip().str.lower().isin(
            ["vacant", "vacancy", "unoccupied", "empty", "available", ""]
        )
    else:
        df["is_vacant"] = df["actual_rent_num"] == 0

    # Revenue gap calculation
    df["monthly_gap"] = df.apply(
        lambda r: r["market_rent_num"] if r["is_vacant"]
                  else max(r["market_rent_num"] - r["actual_rent_num"], 0),
        axis=1
    )
    df["annual_gap"] = df["monthly_gap"] * 12

    # Status label
    df["issue"] = df.apply(
        lambda r: "🔴 VACANT" if r["is_vacant"]
                  else ("🟡 BELOW RATE" if r["monthly_gap"] > 0 else "🟢 OK"),
        axis=1
    )
    return df

def style_cell(cell, bold=False, color=None, bg=None, align="left", size=11, wrap=False):
    cell.font = Font(name="Arial", bold=bold, size=size, color=color or "000000")
    if bg:
        cell.fill = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)

def thin_border():
    s = Side(style="thin", color="CBD5E1")
    return Border(left=s, right=s, top=s, bottom=s)

def write_excel(df, output_path):
    wb = openpyxl.Workbook()

    # ── SHEET 1: SUMMARY DASHBOARD ──────────────────────────────
    ws1 = wb.active
    ws1.title = "📊 Executive Summary"
    ws1.sheet_view.showGridLines = False
    ws1.column_dimensions["A"].width = 3

    # Header banner
    ws1.merge_cells("B2:I2")
    c = ws1["B2"]
    c.value = "PORTFOLIO REVENUE LEAKAGE ANALYSIS"
    style_cell(c, bold=True, color=WHITE, bg=NAVY, align="center", size=16)
    ws1.row_dimensions[2].height = 36

    ws1.merge_cells("B3:I3")
    c = ws1["B3"]
    c.value = f"Generated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}  |  Total Units Analyzed: {len(df)}"
    style_cell(c, color="FFFFFF", bg=ACCENT_BLUE, align="center", size=10)
    ws1.row_dimensions[3].height = 20

    # KPI boxes
    total_units     = len(df)
    vacant_units    = df["is_vacant"].sum()
    below_rate      = ((~df["is_vacant"]) & (df["monthly_gap"] > 0)).sum()
    healthy_units   = total_units - vacant_units - below_rate
    monthly_leakage = df["monthly_gap"].sum()
    annual_leakage  = df["annual_gap"].sum()
    occupancy_rate  = ((total_units - vacant_units) / total_units * 100) if total_units else 0
    current_income  = df[~df["is_vacant"]]["actual_rent_num"].sum()
    potential_income= df["market_rent_num"].sum()

    kpis = [
        ("Total Units",        total_units,              WHITE,       NAVY,        "center"),
        ("Vacant Units",       vacant_units,             WHITE,       RED,         "center"),
        ("Below Market Rate",  below_rate,               "000000",    YELLOW,      "center"),
        ("Healthy Units",      healthy_units,            WHITE,       GREEN,       "center"),
        ("Monthly Leakage",    f"${monthly_leakage:,.0f}", WHITE,     RED,         "center"),
        ("Annual Leakage",     f"${annual_leakage:,.0f}",  WHITE,     "7F1D1D",    "center"),
        ("Occupancy Rate",     f"{occupancy_rate:.1f}%",   WHITE,     ACCENT_BLUE, "center"),
        ("Current Income/Mo",  f"${current_income:,.0f}",  WHITE,     "1E3A5F",    "center"),
    ]

    ws1.row_dimensions[5].height = 28
    ws1.row_dimensions[6].height = 44
    ws1.row_dimensions[7].height = 28

    kpi_cols = ["B", "C", "D", "E", "F", "G", "H", "I"]
    col_widths = [18, 16, 20, 16, 20, 20, 18, 20]
    for i, (col, w) in enumerate(zip(kpi_cols, col_widths)):
        ws1.column_dimensions[col].width = w

    for i, (label, value, fcolor, bg, align) in enumerate(kpis):
        col = kpi_cols[i]
        lc = ws1[f"{col}5"]
        lc.value = label
        style_cell(lc, bold=True, color="94A3B8", bg="1E293B", align="center", size=9)

        vc = ws1[f"{col}6"]
        vc.value = value
        style_cell(vc, bold=True, color=fcolor, bg=bg, align=align, size=14)

    # Section: Revenue breakdown
    ws1.row_dimensions[9].height = 24
    ws1.merge_cells("B9:I9")
    c = ws1["B9"]
    c.value = "  REVENUE BREAKDOWN"
    style_cell(c, bold=True, color=WHITE, bg=NAVY, size=11)

    rev_headers = ["Category", "Units", "Current Monthly", "Potential Monthly", "Monthly Gap", "Annual Gap", "% of Total Gap"]
    rev_data = [
        ["Vacant Units",       vacant_units, 0,            df[df["is_vacant"]]["market_rent_num"].sum(),    df[df["is_vacant"]]["monthly_gap"].sum(),       df[df["is_vacant"]]["annual_gap"].sum()],
        ["Below Market Rate",  below_rate,   df[(~df["is_vacant"])&(df["monthly_gap"]>0)]["actual_rent_num"].sum(), df[(~df["is_vacant"])&(df["monthly_gap"]>0)]["market_rent_num"].sum(), df[(~df["is_vacant"])&(df["monthly_gap"]>0)]["monthly_gap"].sum(), df[(~df["is_vacant"])&(df["monthly_gap"]>0)]["annual_gap"].sum()],
        ["Healthy / On Rate",  healthy_units, df[(~df["is_vacant"])&(df["monthly_gap"]==0)]["actual_rent_num"].sum(), df[(~df["is_vacant"])&(df["monthly_gap"]==0)]["market_rent_num"].sum(), 0, 0],
        ["TOTAL PORTFOLIO",    total_units,  current_income, potential_income, monthly_leakage, annual_leakage],
    ]

    for j, h in enumerate(rev_headers):
        c = ws1.cell(row=10, column=j+2)
        c.value = h
        style_cell(c, bold=True, color=WHITE, bg=ACCENT_BLUE, align="center", size=10)
        ws1.row_dimensions[10].height = 22

    for r, row in enumerate(rev_data):
        is_total = r == len(rev_data) - 1
        bg_color = "1E293B" if is_total else (LIGHT_RED if r == 0 else LIGHT_BLUE if r == 1 else LIGHT_GREEN)
        txt_color = WHITE if is_total else "0F172A"
        row_num = 11 + r
        ws1.row_dimensions[row_num].height = 22

        for j, val in enumerate(row):
            c = ws1.cell(row=row_num, column=j+2)
            if j == 0:
                c.value = val
                style_cell(c, bold=is_total, color=txt_color, bg=bg_color, size=10)
            elif j == 1:
                c.value = val
                style_cell(c, bold=is_total, color=txt_color, bg=bg_color, align="center", size=10)
            else:
                c.value = val
                c.number_format = '$#,##0'
                style_cell(c, bold=is_total, color=txt_color, bg=bg_color, align="right", size=10)
            c.border = thin_border()

        # % of total gap column
        gap_val = row[4]
        pct = (gap_val / monthly_leakage * 100) if monthly_leakage > 0 else 0
        c = ws1.cell(row=row_num, column=8)
        c.value = f"{pct:.1f}%" if not is_total else "100.0%"
        style_cell(c, bold=is_total, color=txt_color, bg=bg_color, align="center", size=10)
        c.border = thin_border()

    # ── SHEET 2: ALL UNITS DETAIL ────────────────────────────────
    ws2 = wb.create_sheet("🏠 All Units")
    ws2.sheet_view.showGridLines = False
    ws2.freeze_panes = "A2"

    headers2 = ["Unit", "Property", "Tenant Name", "Lease Start", "Lease End",
                "Market Rent", "Actual Rent", "Monthly Gap", "Annual Gap", "Status"]
    col_w2 = [12, 20, 24, 14, 14, 14, 14, 14, 14, 16]
    for i, (h, w) in enumerate(zip(headers2, col_w2)):
        c = ws2.cell(row=1, column=i+1)
        c.value = h
        style_cell(c, bold=True, color=WHITE, bg=NAVY, align="center", size=10)
        ws2.column_dimensions[get_column_letter(i+1)].width = w
    ws2.row_dimensions[1].height = 24

    sort_order = {"🔴 VACANT": 0, "🟡 BELOW RATE": 1, "🟢 OK": 2}
    df_sorted = df.sort_values("issue", key=lambda x: x.map(sort_order))

    for r, (_, row) in enumerate(df_sorted.iterrows()):
        rn = r + 2
        if row["issue"] == "🔴 VACANT":
            bg = LIGHT_RED
        elif row["issue"] == "🟡 BELOW RATE":
            bg = YELLOW
        else:
            bg = WHITE if r % 2 == 0 else GRAY

        vals = [
            row.get("unit", ""),
            row.get("property", ""),
            row.get("tenant_name", "VACANT"),
            row.get("lease_start", ""),
            row.get("lease_end", ""),
            row["market_rent_num"],
            row["actual_rent_num"],
            row["monthly_gap"],
            row["annual_gap"],
            row["issue"],
        ]

        for j, val in enumerate(vals):
            c = ws2.cell(row=rn, column=j+1)
            c.value = val
            is_currency = j in [5, 6, 7, 8]
            c.alignment = Alignment(horizontal="center" if j in [0,3,4,9] else "left", vertical="center")
            c.font = Font(name="Arial", size=10)
            c.fill = PatternFill("solid", start_color=bg)
            c.border = thin_border()
            if is_currency and val != "":
                c.number_format = '$#,##0'
        ws2.row_dimensions[rn].height = 18

    # ── SHEET 3: VACANCIES ONLY ──────────────────────────────────
    ws3 = wb.create_sheet("🔴 Vacancies")
    ws3.sheet_view.showGridLines = False

    ws3.merge_cells("A1:G1")
    c = ws3["A1"]
    c.value = f"VACANT UNITS  —  {vacant_units} units  |  ${df[df['is_vacant']]['monthly_gap'].sum():,.0f}/mo lost  |  ${df[df['is_vacant']]['annual_gap'].sum():,.0f}/yr lost"
    style_cell(c, bold=True, color=WHITE, bg=RED, align="center", size=12)
    ws3.row_dimensions[1].height = 28

    headers3 = ["Unit", "Property", "Market Rent", "Monthly Lost", "Annual Lost", "Last Lease End", "Days Vacant Est."]
    col_w3 = [12, 22, 14, 14, 14, 18, 18]
    for i, (h, w) in enumerate(zip(headers3, col_w3)):
        c = ws3.cell(row=2, column=i+1)
        c.value = h
        style_cell(c, bold=True, color=WHITE, bg="7F1D1D", align="center", size=10)
        ws3.column_dimensions[get_column_letter(i+1)].width = w

    vacant_df = df[df["is_vacant"]].sort_values("market_rent_num", ascending=False)
    for r, (_, row) in enumerate(vacant_df.iterrows()):
        rn = r + 3
        bg = LIGHT_RED if r % 2 == 0 else "FFF1F1"

        lease_end = row.get("lease_end", "")
        days_vacant = ""
        if lease_end and str(lease_end).strip() not in ["", "nan", "NaT"]:
            try:
                end_dt = pd.to_datetime(lease_end)
                days_vacant = (datetime.now() - end_dt.to_pydatetime()).days
                days_vacant = max(days_vacant, 0)
            except:
                days_vacant = ""

        for j, val in enumerate([
            row.get("unit", ""),
            row.get("property", ""),
            row["market_rent_num"],
            row["monthly_gap"],
            row["annual_gap"],
            lease_end,
            days_vacant
        ]):
            c = ws3.cell(row=rn, column=j+1)
            c.value = val
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.font = Font(name="Arial", size=10)
            c.fill = PatternFill("solid", start_color=bg)
            c.border = thin_border()
            if j in [2, 3, 4]:
                c.number_format = '$#,##0'
        ws3.row_dimensions[rn].height = 18

    # ── SHEET 4: BELOW RATE ──────────────────────────────────────
    ws4 = wb.create_sheet("🟡 Below Market Rate")
    ws4.sheet_view.showGridLines = False

    below_df = df[(~df["is_vacant"]) & (df["monthly_gap"] > 0)].sort_values("monthly_gap", ascending=False)

    ws4.merge_cells("A1:H1")
    c = ws4["A1"]
    c.value = f"BELOW MARKET RATE  —  {len(below_df)} tenants  |  ${below_df['monthly_gap'].sum():,.0f}/mo under market  |  ${below_df['annual_gap'].sum():,.0f}/yr opportunity"
    style_cell(c, bold=True, color="000000", bg=YELLOW, align="center", size=12)
    ws4.row_dimensions[1].height = 28

    headers4 = ["Unit", "Tenant Name", "Lease Start", "Lease End", "Market Rent", "Actual Rent", "Monthly Gap", "Annual Opportunity"]
    col_w4 = [12, 24, 14, 14, 14, 14, 14, 18]
    for i, (h, w) in enumerate(zip(headers4, col_w4)):
        c = ws4.cell(row=2, column=i+1)
        c.value = h
        style_cell(c, bold=True, color=WHITE, bg="92400E", align="center", size=10)
        ws4.column_dimensions[get_column_letter(i+1)].width = w

    for r, (_, row) in enumerate(below_df.iterrows()):
        rn = r + 3
        bg = YELLOW if r % 2 == 0 else "FFFBEB"
        for j, val in enumerate([
            row.get("unit", ""),
            row.get("tenant_name", ""),
            row.get("lease_start", ""),
            row.get("lease_end", ""),
            row["market_rent_num"],
            row["actual_rent_num"],
            row["monthly_gap"],
            row["annual_gap"],
        ]):
            c = ws4.cell(row=rn, column=j+1)
            c.value = val
            c.alignment = Alignment(horizontal="center" if j in [0,2,3] else "left", vertical="center")
            c.font = Font(name="Arial", size=10)
            c.fill = PatternFill("solid", start_color=bg)
            c.border = thin_border()
            if j in [4, 5, 6, 7]:
                c.number_format = '$#,##0'
        ws4.row_dimensions[rn].height = 18

    wb.save(output_path)
    print(f"\n✅ Analysis complete: {output_path}")
    return monthly_leakage, annual_leakage, vacant_units, below_rate

# ─────────────────────────────────────────────
#  MAIN
# ─────────────────────────────────────────────
def main():
    # Check if a file was passed as argument
    input_file = sys.argv[1] if len(sys.argv) > 1 else INPUT_FILE
    output_file = sys.argv[2] if len(sys.argv) > 2 else OUTPUT_FILE

    if not os.path.exists(input_file):
        print(f"\n⚠️  File not found: {input_file}")
        print("Please update INPUT_FILE in the script or pass your file as an argument:")
        print("  python rent_analyzer.py your_export.csv")
        print("\n📋 GENERATING SAMPLE DEMO OUTPUT instead...\n")
        df = generate_sample_data()
    else:
        print(f"📂 Loading: {input_file}")
        df = load_data(input_file)
        df = normalize_columns(df)

    df = analyze(df)
    monthly, annual, vacancies, below = write_excel(df, output_file)

    print(f"\n{'='*52}")
    print(f"  📊 PORTFOLIO REVENUE ANALYSIS COMPLETE")
    print(f"{'='*52}")
    print(f"  🔴 Vacant Units:          {vacancies}")
    print(f"  🟡 Below Market Rate:     {below}")
    print(f"  💸 Monthly Revenue Gap:   ${monthly:,.0f}")
    print(f"  💰 Annual Revenue Gap:    ${annual:,.0f}")
    print(f"{'='*52}")
    print(f"  📁 Output: {output_file}\n")

def generate_sample_data():
    """Generate realistic sample data for demo purposes."""
    import random
    random.seed(42)
    units = []
    properties = ["Elm Street Complex", "Oak Avenue", "Maple Gardens", "Pine Ridge"]

    for i in range(1, 51):
        prop = properties[i % len(properties)]
        unit = f"{prop[:3].upper()}-{i:03d}"
        is_vacant = random.random() < 0.10
        market = random.choice([950, 1050, 1100, 1200, 1350, 1450, 1550, 1650])
        
        if is_vacant:
            status = "Vacant"
            actual = 0
            tenant = ""
            lease_end = pd.Timestamp("2024-11-01") + pd.Timedelta(days=random.randint(0, 120))
        else:
            status = "Occupied"
            discount = random.choice([0, 0, 0, 25, 50, 75, 100, 150])
            actual = market - discount
            tenant = random.choice([
                "Smith, John", "Garcia, Maria", "Johnson, David", "Williams, Sarah",
                "Brown, Michael", "Davis, Jennifer", "Miller, Robert", "Wilson, Lisa",
                "Moore, James", "Taylor, Patricia", "Anderson, Charles", "Thomas, Barbara"
            ])
            lease_end = pd.Timestamp("2025-06-01") + pd.Timedelta(days=random.randint(0, 365))

        units.append({
            "Unit": unit,
            "Property": prop,
            "Tenant Name": tenant,
            "Lease Start": (pd.Timestamp("2023-01-01") + pd.Timedelta(days=random.randint(0, 365))).strftime("%Y-%m-%d"),
            "Lease End": lease_end.strftime("%Y-%m-%d"),
            "Market Rent": market,
            "Actual Rent": actual,
            "Status": status,
        })

    df = pd.DataFrame(units)
    df = normalize_columns(df)
    return df

if __name__ == "__main__":
    main()
